"""Testdateien: kleine echte PDF/DOCX/PPTX und ein State, der auf sie zeigt."""
import io
from pathlib import Path


def make_pdf(pages: list[str]) -> bytes:
    """Minimales, gültiges PDF mit einer Textzeile je Seite (WinAnsi, damit Umlaute und ß stimmen)."""
    n = len(pages)
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{' '.join(f'{4 + 2 * i} 0 R' for i in range(n))}] /Count {n} >>".encode(),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"]
    for i, text in enumerate(pages):
        esc = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 50 700 Td ({esc}) Tj ET".encode("cp1252")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>".encode())
        objs.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for o in offsets:
        out += f"{o:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_encrypted_pdf(pages: list[str], password: str = "geheim") -> bytes:
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(pages))))
    w.encrypt(password)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def make_docx(paragraphs: list[str]) -> bytes:
    from docx import Document
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx(slides: list[str]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for text in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def put(ziel: Path, rel: str, data: bytes) -> Path:
    p = ziel / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)
    return p


def make_state(items, course_id: int = 7, name: str = "Sportpsychologie", dir_: str = "Kurs") -> dict:
    """items: [(cmid, typ, name, [relative Dateien])] -> State im Format von sync_state.json (ein Kurs, ein Abschnitt)."""
    cms = {str(cmid): {"name": n, "typ": typ, "section": "01 Woche 1", "url": f"https://x/mod/{typ}/view.php?id={cmid}",
                       "files": files, "skipped": []} for cmid, typ, n, files in items}
    return {"last_run": 1, "digest": [],
            "courses": {str(course_id): {"id": course_id, "name": name, "url": f"https://x/course/view.php?id={course_id}",
                                         "dir": dir_, "layout": [["Woche 1", list(cms)]], "cms": cms}}}
